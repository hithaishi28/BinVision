from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path

OUT = Path(__file__).parent / 'BinVision_Impact_Report.docx'
doc = Document()
sec = doc.sections[0]
sec.page_width = Inches(8.27)
sec.page_height = Inches(11.69)
sec.top_margin = Inches(.50)
sec.bottom_margin = Inches(.50)
sec.left_margin = Inches(.62)
sec.right_margin = Inches(.62)

navy = '132338'; green = '14846B'; teal = '20A99A'; muted = '56697C'; pale = 'EAF4F2'; pale2 = 'F3F6F8'; amber='9B6B10'
styles = doc.styles
normal = styles['Normal']
normal.font.name = 'Aptos'; normal.font.size = Pt(9.15); normal.font.color.rgb = RGBColor.from_string('2D3D4E')
normal.paragraph_format.space_after = Pt(3.2); normal.paragraph_format.line_spacing = 1.06
for name, size, color in [('Title', 24, navy), ('Heading 1', 12, green), ('Heading 2', 9.5, navy)]:
    st=styles[name]; st.font.name='Aptos Display' if name=='Title' else 'Aptos'; st.font.size=Pt(size); st.font.bold=True; st.font.color.rgb=RGBColor.from_string(color)
    st.paragraph_format.space_before=Pt(5); st.paragraph_format.space_after=Pt(2.5); st.paragraph_format.keep_with_next=True

def shade(cell, fill):
    tcPr=cell._tc.get_or_add_tcPr(); shd=OxmlElement('w:shd'); shd.set(qn('w:fill'),fill); tcPr.append(shd)

def set_cell_margins(cell, top=90, start=120, bottom=90, end=120):
    tc=cell._tc; tcPr=tc.get_or_add_tcPr(); mar=tcPr.first_child_found_in('w:tcMar')
    if mar is None: mar=OxmlElement('w:tcMar'); tcPr.append(mar)
    for key,val in [('top',top),('start',start),('bottom',bottom),('end',end)]:
        node=mar.find(qn('w:'+key))
        if node is None: node=OxmlElement('w:'+key); mar.append(node)
        node.set(qn('w:w'),str(val)); node.set(qn('w:type'),'dxa')

def run(p, text, size=None, bold=False, color=None, italic=False):
    r=p.add_run(text); r.font.name='Aptos'; r.font.size=Pt(size or 9.15); r.bold=bold; r.italic=italic
    if color: r.font.color.rgb=RGBColor.from_string(color)
    return r

# Title and assessment line
p=doc.add_paragraph(style='Title'); p.paragraph_format.space_after=Pt(0); run(p,'BinVision Impact Report',24,True,navy)
p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(7)
run(p,'IMPLEMENTATION SUSTAINABILITY  /  PS-1B  /  CLIMATE & ENVIRONMENT',8.3,True,green)

# Executive assessment callout
t=doc.add_table(rows=1,cols=1); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False; t.columns[0].width=Inches(7.0)
c=t.cell(0,0); shade(c,pale); set_cell_margins(c,130,180,130,180); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
p=c.paragraphs[0]; p.paragraph_format.space_after=Pt(0); p.paragraph_format.line_spacing=1.05
run(p,'ASSESSMENT  ',9,True,green); run(p,'Promising as a software implementation, with environmental value dependent on validated detection and real user adoption. The current demo demonstrates the workflow, not measured waste outcomes.',9.1,False,navy)

def heading(text): doc.add_paragraph(text,style='Heading 1')

heading('Why the implementation is practical')
doc.add_paragraph('BinVision is software-first. Its existing React and TypeScript interface, shared detection response, FastAPI boundaries, YOLO model adapter, and Groq-backed assistant form a modular product path. The stable response format lets the model integration evolve without redesigning the user experience. The front end already supports upload, detection presentation, category summaries, local session analytics, and context-aware assistant interactions. This is a credible prototype foundation that can be piloted incrementally.')

heading('Where environmental benefit could come from')
doc.add_paragraph('Image-based explanations may help people recognize recyclable, organic, and hazardous items and make more informed sorting choices. Clearer identification could support cleaner material streams and safer handling decisions. The assistant can explain a result and point users toward responsible disposal guidance. These are plausible contribution pathways. BinVision has not yet demonstrated that users sort more accurately, that contamination falls, or that any material is recovered as a result.')

heading('Conditions for responsible scale')
doc.add_paragraph('The current front end defaults to mock detections, and its analytics are local demo data. Before an impact claim or operational rollout, connect the intended model and validate performance on representative images for each category, including confidence calibration and common misclassifications. Review category mapping and hazardous-item guidance with qualified local sources. Keep the assistant grounded in returned detections, show uncertainty plainly, and direct users to local instructions where disposal rules vary. Minimize image retention and access. Track inference and storage needs so cloud use remains proportionate to the service provided.')

heading('Measure the outcome, not just activity')
doc.add_paragraph('A pilot should separate system performance from environmental outcomes. Record detection precision and recall by category on reviewed data; monitor user corrections and repeat use; then work with a participating site to compare sorting quality before and after guidance using a consistent audit method. Measure material recovery or contamination only where a partner can provide direct, documented measurements. Do not treat images analyzed, mock detections, confidence values, or dashboard counts as tonnes diverted.')

# Three compact measurement priorities
tbl=doc.add_table(rows=1,cols=3); tbl.alignment=WD_TABLE_ALIGNMENT.CENTER; tbl.autofit=False
widths=[2.28,2.28,2.44]
for i,w in enumerate(widths): tbl.columns[i].width=Inches(w)
for i,(a,b) in enumerate([
    ('MODEL QUALITY','Reviewed per-category precision, recall, and calibration'),
    ('USER VALUE','Corrections, guidance use, and repeat engagement'),
    ('FIELD IMPACT','Audited sorting quality and partner-measured recovery')]):
    c=tbl.cell(0,i); c.width=Inches(widths[i]); shade(c,pale2); set_cell_margins(c,100,120,100,120)
    p=c.paragraphs[0]; p.paragraph_format.space_after=Pt(2); run(p,a,7.6,True,green)
    p=c.add_paragraph(); p.paragraph_format.space_after=Pt(0); p.paragraph_format.line_spacing=1.02; run(p,b,8.0,False,navy)

heading('Overall outlook')
p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(0)
run(p,'BinVision is a feasible, potentially useful digital sustainability tool. ',9.15,True,navy)
run(p,'Its strongest next step is a bounded pilot with validated model output, locally reviewed guidance, privacy-aware data handling, and an agreed measurement plan. Evidence from that pilot should determine whether to scale.',9.15,False)

# Footer label without decorative rules
footer=sec.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT; footer.paragraph_format.space_before=Pt(0)
run(footer,'BINVISION  ·  QUALITATIVE IMPLEMENTATION ASSESSMENT',7.2,True,'8493A2')

doc.save(OUT)
print(f'Created {OUT}')
