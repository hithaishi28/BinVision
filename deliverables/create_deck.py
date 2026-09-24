from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.oxml import parse_xml
from pptx.oxml.ns import nsdecls
from docx import Document
from docx.shared import Inches as DInches, Pt as DPt, RGBColor as DRGB
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path

OUT = Path(__file__).parent
W, H = 13.333, 7.5
BG = '0B1220'; PANEL = '111C2D'; PANEL2 = '152437'; WHITE = 'F5F8FC'; MUTED = '9DAEC2'
GREEN = '45D6A1'; TEAL = '38C5C0'; LIME = 'B7E36A'; AMBER = 'F3BE5E'; RED = 'FF727C'; BLUE = '7AA7FF'
FONT = 'Aptos'

def rgb(h): return RGBColor.from_string(h)

prs = Presentation()
prs.slide_width = Inches(W); prs.slide_height = Inches(H)
blank = prs.slide_layouts[6]

def shape(slide, x,y,w,h, fill=PANEL, line=None, radius=True, transparency=0):
    t = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    s = slide.shapes.add_shape(t, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid(); s.fill.fore_color.rgb=rgb(fill)
    if transparency: s.fill.transparency=transparency
    s.line.fill.background() if not line else None
    if line:
        s.line.color.rgb=rgb(line); s.line.width=Pt(1)
    if radius:
        try: s.adjustments[0] = 0.12
        except Exception: pass
    return s

def line(slide,x1,y1,x2,y2,color=GREEN,width=1.5,dash=None,begin=None,end=None):
    s=slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    s.line.color.rgb=rgb(color); s.line.width=Pt(width)
    if dash: s.line.dash_style=MSO_LINE_DASH_STYLE.DASH
    if begin: s.line.begin_arrowhead=begin
    if end: s.line.end_arrowhead=end
    return s

def txt(slide,text,x,y,w,h,size=16,color=WHITE,bold=False,font=FONT,align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.MIDDLE, margin=0.02, italic=False):
    box=slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf=box.text_frame; tf.clear(); tf.word_wrap=True
    tf.margin_left=Inches(margin); tf.margin_right=Inches(margin); tf.margin_top=Inches(margin); tf.margin_bottom=Inches(margin); tf.vertical_anchor=valign
    p=tf.paragraphs[0]; p.alignment=align
    r=p.add_run(); r.text=text; r.font.name=font; r.font.size=Pt(size); r.font.bold=bold; r.font.italic=italic; r.font.color.rgb=rgb(color)
    return box

def rich(slide,runs,x,y,w,h,size=16,align=PP_ALIGN.LEFT):
    box=slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)); tf=box.text_frame; tf.clear(); tf.word_wrap=True
    tf.margin_left=tf.margin_right=Inches(.02); tf.margin_top=tf.margin_bottom=Inches(.02)
    p=tf.paragraphs[0]; p.alignment=align
    for text,color,bold in runs:
        r=p.add_run(); r.text=text; r.font.name=FONT; r.font.size=Pt(size); r.font.bold=bold; r.font.color.rgb=rgb(color)
    return box

def base(title,kicker=None,n=None):
    s=prs.slides.add_slide(blank); s.background.fill.solid(); s.background.fill.fore_color.rgb=rgb(BG)
    # restrained corner glow motifs
    shape(s,12.78,-.16,.72,.72,'112A35',None,True,5)
    if kicker: txt(s,kicker.upper(),.62,.36,8,.24,9,GREEN,True)
    if title: txt(s,title,.62,.70,11.9,.62,27,WHITE,True)
    line(s,.64,7.12,12.68,7.12,'26364A',.7)
    txt(s,'BINVISION   /   PS-1B',.64,7.17,4,.16,8,MUTED,True)
    if n: txt(s,f'{n:02d} / 10',11.8,7.15,.85,.18,8,MUTED,True,align=PP_ALIGN.RIGHT)
    return s

def pill(slide,label,x,y,w,color=GREEN,bg=PANEL2):
    shape(slide,x,y,w,.31,bg,None,True)
    shape(slide,x+.10,y+.105,.09,.09,color,None,True)
    txt(slide,label,x+.25,y+.025,w-.3,.25,9,color,True)

def card(slide,x,y,w,h,title,body=None,accent=GREEN,tag=None):
    shape(slide,x,y,w,h,PANEL,'223249',True)
    shape(slide,x,y,.055,h,accent,None,False)
    if tag: txt(slide,tag,x+.22,y+.16,w-.4,.2,9,accent,True)
    txt(slide,title,x+.22,y+(.34 if tag else .2),w-.42,.36,16,WHITE,True)
    if body: txt(slide,body,x+.22,y+(.79 if tag else .67),w-.42,h-(1.0 if tag else .83),11,MUTED,False,valign=MSO_ANCHOR.TOP)

def node(slide,x,y,w,label,sub,color=GREEN):
    shape(slide,x,y,w,.84,PANEL2,'29435A',True)
    shape(slide,x+.14,y+.24,.34,.34,color,None,True)
    txt(slide,label,x+.60,y+.10,w-.72,.29,12,WHITE,True)
    if sub: txt(slide,sub,x+.60,y+.42,w-.72,.24,9,MUTED)

def bubble(slide,x,y,w,h,text,who='assistant',accent=GREEN):
    shape(slide,x,y,w,h,PANEL2 if who=='assistant' else '18352F',None,True)
    txt(slide,text,x+.15,y+.08,w-.3,h-.16,11,WHITE,False,valign=MSO_ANCHOR.MIDDLE)

# 01 — Title
s=prs.slides.add_slide(blank); s.background.fill.solid(); s.background.fill.fore_color.rgb=rgb(BG)
shape(s,9.8,.3,3.8,3.8,'102634',None,True,25)
txt(s,'BINVISION',.72,.76,7.2,.82,34,WHITE,True)
txt(s,'AI-Powered Waste Intelligence\n& Classification',.76,1.75,6.4,1.18,23,WHITE,False)
pill(s,'PS-1B  /  AI-DRIVEN WASTE CLASSIFICATION',.76,3.28,3.65,GREEN)
pill(s,'CLIMATE & ENVIRONMENT',.76,3.72,2.35,TEAL)
txt(s,'Team: [TEAM NAME]   ·   [MEMBER 1]   ·   [MEMBER 2]   ·   [MEMBER 3]',.78,6.55,8,.26,10,MUTED)
# hero flow
node(s,7.2,1.33,1.45,'WASTE','image input',BLUE)
line(s,8.72,1.75,9.16,1.75,TEAL,2)
node(s,9.19,1.33,1.55,'VISION','detect + classify',TEAL)
line(s,10.80,1.75,11.17,1.75,GREEN,2)
shape(s,11.20,1.05,1.25,1.46,PANEL,'29435A',True)
for i,(lab,col) in enumerate([('RECYCLE',GREEN),('ORGANIC',AMBER),('HAZARD',RED)]):
    shape(s,11.34,1.19+i*.41,.16,.16,col,None,True); txt(s,lab,11.60,1.16+i*.41,.73,.2,7,col,True)
txt(s,'See waste. Understand waste. Sort smarter.',7.15,2.28,5.2,.38,13,WHITE,True)
for i,(lab,col) in enumerate([('RECYCLABLE',GREEN),('ORGANIC',AMBER),('HAZARDOUS',RED)]): pill(s,lab,7.2+i*1.72,3.10,1.55,col)
shape(s,7.2,4.12,5.25,1.48,PANEL,'223249',True)
txt(s,'CLASS',7.48,4.4,1,.22,8,MUTED,True); txt(s,'CONFIDENCE',9.05,4.4,1.2,.22,8,MUTED,True); txt(s,'BOX',11.0,4.4,.65,.22,8,MUTED,True)
txt(s,'Structured detections, ready for action',7.48,4.82,4.4,.36,15,WHITE,True)
txt(s,'Software intelligence for better sorting decisions',.78,5.48,6.2,.34,13,GREEN,True)
txt(s,'01 / 10',11.8,7.15,.85,.18,8,MUTED,True,align=PP_ALIGN.RIGHT)

# 02 — Problem
s=base('Mixed waste creates an identification gap','THE CHALLENGE',2)
txt(s,'Sorting decisions are made with incomplete visual guidance.',.68,1.5,11.4,.38,17,WHITE,True)
items=[('Mixed materials','One image can contain several waste types.',GREEN),('Contamination risk','Misidentification can compromise recyclable material.',AMBER),('Hazard uncertainty','Batteries and chemical containers need careful handling.',RED),('Inconsistent decisions','Manual identification varies across people and moments.',BLUE)]
for i,(t,b,c) in enumerate(items):
    x=.68+(i%2)*3.08; y=2.12+(i//2)*1.33
    card(s,x,y,2.82,1.12,t,b,c)
# visual chain
labels=[('MIXED WASTE',GREEN),('IDENTIFICATION\nCHALLENGE',BLUE),('SORTING ERROR',AMBER),('MATERIAL\nRECOVERY AT RISK',RED)]
for i,(lab,c) in enumerate(labels):
    x=.70+i*3.12; shape(s,x,5.24,2.48,.91,PANEL2,'29435A',True)
    shape(s,x+.16,5.51,.28,.28,c,None,True); txt(s,lab,x+.56,5.34,1.75,.66,11,WHITE,True)
    if i<3: line(s,x+2.55,5.70,x+3.01,5.70,MUTED,1.3)
txt(s,'A clear visual explanation can help people make more informed sorting choices.',.72,6.48,10.8,.28,11,MUTED,False,italic=True)

# 03 — Solution
s=base('BinVision turns an image into an explainable result','THE SOLUTION',3)
txt(s,'A software intelligence layer for smarter waste-sorting decisions.',.68,1.42,11,.34,16,WHITE,True)
steps=[('01','Waste image','Upload'),('02','AI vision analysis','Detect objects'),('03','Category mapping','Three outputs'),('04','Confidence + box','Show evidence'),('05','Results & analytics','Summarize'),('06','AI assistant','Explain')]
for i,(num,t,sub) in enumerate(steps):
    x=.68+i*2.06
    shape(s,x,2.20,1.78,1.36,PANEL,'223249',True)
    txt(s,num,x+.16,2.37,.43,.23,9,GREEN,True)
    txt(s,t,x+.16,2.72,1.46,.43,12,WHITE,True)
    txt(s,sub,x+.16,3.20,1.42,.2,9,MUTED)
    if i<5: line(s,x+1.82,2.88,x+2.00,2.88,TEAL,1.5)
for i,(label,c) in enumerate([('Recyclable',GREEN),('Organic',AMBER),('Hazardous',RED)]):
    pill(s,label,1.05+i*2.0,4.22,1.65,c)
card(s,7.22,4.08,5.25,1.55,'One stable response shape','Class  ·  Category  ·  Confidence  ·  Bounding box',TEAL,'FRONTEND CONTRACT')
txt(s,'Model internals can evolve while the product UI keeps a consistent detection format.',.72,6.10,11.5,.38,12,MUTED)

# 04 — Journey
s=base('From upload to useful guidance','USER JOURNEY',4)
journey=[('Upload','Image input'),('Preview','Check image'),('Analyze','Start request'),('Detect','Find objects'),('Classify','Assign category'),('Review','Read results'),('Ask AI','Get guidance')]
for i,(a,b) in enumerate(journey):
    x=.66+i*1.79
    shape(s,x,1.67,1.52,1.02,PANEL2,'29435A',True)
    txt(s,f'{i+1:02}',x+.14,1.79,.36,.22,9,GREEN,True)
    txt(s,a,x+.14,2.05,1.26,.27,12,WHITE,True)
    txt(s,b,x+.14,2.38,1.25,.18,8,MUTED)
    if i<6: line(s,x+1.55,2.18,x+1.73,2.18,TEAL,1.3)
# detection UI panel
shape(s,.68,3.17,7.42,3.24,PANEL,'223249',True)
txt(s,'AI DETECTION PREVIEW',.96,3.40,3.2,.25,10,MUTED,True)
shape(s,.96,3.82,4.18,2.06,'1C2B37',None,True)
shape(s,2.36,4.08,.82,1.40,'405E55',None,True)
shape(s,2.46,3.91,.61,.22,'C6D7CF',None,True)
shape(s,2.15,3.99,1.22,1.64,GREEN,None,False)
shape(s,1.23,4.75,.63,.62,'B68243',None,True)
shape(s,3.76,4.82,.88,.54,'76879A',None,True)
pill(s,'Plastic bottle  ·  demo bbox',1.17,5.94,2.7,GREEN)
txt(s,'DEMO SUMMARY',5.45,3.60,2.2,.22,9,MUTED,True)
for i,(label,val,col) in enumerate([('Detected items','6',WHITE),('Avg confidence','89%',TEAL),('Recyclable','3',GREEN),('Organic','2',AMBER),('Hazardous','1',RED)]):
    yy=3.98+i*.43; txt(s,label,5.48,yy,1.45,.24,10,MUTED); txt(s,val,7.15,yy,0.58,.24,11,col,True,align=PP_ALIGN.RIGHT)
shape(s,8.43,3.17,4.02,3.24,PANEL,'223249',True)
txt(s,'RESULT SIGNALS',8.74,3.43,3.3,.24,10,MUTED,True)
for i,(lab,col) in enumerate([('Recyclable',GREEN),('Organic',AMBER),('Hazardous',RED)]):
    shape(s,8.77,3.92+i*.70,.28,.28,col,None,True); txt(s,lab,9.20,3.91+i*.70,1.6,.28,13,WHITE,True)
    txt(s,['count','count','count'][i],11.05,3.92+i*.70,.8,.22,9,MUTED,align=PP_ALIGN.RIGHT)
txt(s,'Confidence + bounding box + category count',8.78,6.00,3.15,.19,9,MUTED)

# 05 — Experience
s=base('A product flow built around the user','PRODUCT EXPERIENCE',5)
txt(s,'Dashboard  ·  Detection  ·  Results  ·  Analytics',.68,1.4,11,.30,14,GREEN,True)
# dashboard mock
shape(s,.68,1.95,3.84,4.45,'EAF1F5',None,True)
shape(s,.68,1.95,3.84,.50,'152336',None,True)
txt(s,'BinVision',.88,2.08,1.3,.22,11,WHITE,True)
txt(s,'Dashboard',.95,2.70,2.3,.27,15,'132338',True)
txt(s,'DEMO VALUES  ·  LOCAL ANALYTICS',.95,2.99,2.85,.15,7,'71849A',True)
for i,(a,v,c) in enumerate([('Objects','24',GREEN),('Recyclable','12',GREEN),('Organic','8',AMBER),('Hazardous','4',RED)]):
    x=.92+(i%2)*1.62; y=3.18+(i//2)*.88
    shape(s,x,y,1.42,.68,'FFFFFF','DCE5EC',True)
    txt(s,a,x+.10,y+.09,1.2,.15,8,'5D7085'); txt(s,v,x+.10,y+.31,.65,.23,14,c,True)
txt(s,'Recent sessions',.95,5.08,2.4,.2,10,'132338',True)
for i in range(2):
    shape(s,.94,5.43+i*.41,3.18,.28,'DCE7EC',None,True)
    txt(s,['Kitchen mix · Demo','Campus sample · Demo'][i],1.08,5.46+i*.41,2.8,.18,8,'55687D')
# main detection mock
shape(s,4.80,1.95,4.24,4.45,'111C2D','29435A',True)
txt(s,'Detection workspace',5.08,2.20,3.6,.3,15,WHITE,True)
shape(s,5.08,2.72,3.67,2.28,'1B2B38',None,True)
shape(s,6.22,3.12,.91,1.48,'405E55',None,True); shape(s,6.32,2.96,.7,.21,'CDD9D4',None,True)
shape(s,6.06,3.00,1.24,1.76,GREEN,None,False)
shape(s,7.33,3.75,.74,.62,'9D744A',None,True)
shape(s,5.41,4.12,.56,.52,'63798B',None,True)
txt(s,'AI box view',5.28,5.15,1.25,.2,9,MUTED)
pill(s,'Analyze image',7.25,5.10,1.38,GREEN)
txt(s,'Original / AI Detection toggle',5.10,5.80,3.5,.25,10,MUTED)
# results mock
shape(s,9.32,1.95,3.12,4.45,'EAF1F5',None,True)
txt(s,'Results',9.58,2.20,2.2,.30,15,'132338',True)
txt(s,'6 detected items',9.60,2.61,2.2,.22,10,'5D7085')
for i,(item,cat,c) in enumerate([('Plastic bottle','Recyclable',GREEN),('Banana peel','Organic',AMBER),('Battery','Hazardous',RED)]):
    y=3.10+i*.83; shape(s,9.58,y,2.55,.67,'FFFFFF','DCE5EC',True)
    txt(s,item,9.75,y+.08,2.12,.20,10,'132338',True); pill(s,cat,9.75,y+.36,1.12,c,'EDF3F5')
txt(s,'Demo values shown',9.62,5.80,2.2,.20,8,'74849A',False,italic=True)

# 06 — Vision intelligence
s=base('Every detection carries context','AI WASTE INTELLIGENCE',6)
txt(s,'Example only  ·  Plastic bottle  ·  Recyclable  ·  94% confidence',.68,1.41,11.4,.30,13,TEAL,True)
shape(s,.68,1.95,5.52,4.55,'172534','29435A',True)
shape(s,1.35,2.44,4.18,3.05,'34424B',None,True)
# illustrative recyclable image scene
shape(s,2.95,2.94,.99,2.03,'447262',None,True); shape(s,3.12,2.75,.65,.28,'C7D7D0',None,True)
shape(s,2.77,2.83,1.35,2.29,GREEN,None,False)
txt(s,'plastic bottle',2.79,2.59,1.25,.18,9,GREEN,True)
shape(s,4.46,3.72,.58,.42,'C48A51',None,True)
txt(s,'Illustrative detection overlay',1.05,5.83,4.5,.22,9,MUTED,False,italic=True)
for i,(lab,val,c) in enumerate([('Class','Plastic bottle',WHITE),('Category','Recyclable',GREEN),('Confidence','94%',TEAL),('Bounding box','[x1, y1, x2, y2]',BLUE)]):
    yy=2.14+i*.72; txt(s,lab,6.62,yy,1.36,.25,10,MUTED,True); txt(s,val,8.12,yy,3.52,.26,13,c,True)
line(s,6.60,5.18,12.22,5.18,'29435A',.8)
txt(s,'YOLO / Vision model',6.63,5.38,1.28,.43,9,WHITE,True,align=PP_ALIGN.CENTER)
txt(s,'Model adapter',8.27,5.38,1.18,.43,10,WHITE,True,align=PP_ALIGN.CENTER)
txt(s,'BinVision UI',10.90,5.38,1.18,.43,11,WHITE,True,align=PP_ALIGN.CENTER)
line(s,8.00,5.58,8.20,5.58,TEAL,1.2); line(s,9.54,5.58,10.82,5.58,TEAL,1.2)
shape(s,6.63,6.10,5.60,.40,'2A2417',None,True)
txt(s,'Frontend defaults to mock data; final model output must be validated.',6.80,6.18,5.28,.20,9,AMBER,True)

# 07 — Assistant
s=base('Ask your waste data','BINVISION AI ASSISTANT',7)
txt(s,'Context-aware explanations and practical sorting guidance.',.68,1.42,11,.30,15,WHITE,True)
shape(s,.68,1.94,6.33,4.68,'121E2E','29435A',True)
shape(s,.68,1.94,6.33,.57,'17273A',None,True)
shape(s,.93,2.11,.22,.22,GREEN,None,True); txt(s,'BinVision AI',1.28,2.08,2.5,.25,12,WHITE,True); pill(s,'DETECTION CONTEXT',4.65,2.07,1.85,TEAL)
bubble(s,.98,2.82,3.60,.64,'Why is this item classified as hazardous?','user')
bubble(s,.98,3.66,5.55,1.18,'This detection is labeled hazardous in the current analysis. Check the item label and local collection guidance before disposing of it.','assistant')
for i,(label,x,width) in enumerate([('Explain results',1.00,1.47),('How to dispose?',2.60,1.55),('What is confidence?',4.28,1.98)]):
    pill(s,label,x,5.29,width,GREEN)
shape(s,1.00,5.91,5.60,.46,'0D1725','29435A',True); txt(s,'Ask about this analysis…',1.20,6.01,3.8,.20,9,MUTED)
shape(s,7.38,1.94,5.05,4.68,PANEL,'223249',True)
txt(s,'ASSISTANT CAPABILITIES',7.70,2.25,4.2,.24,9,GREEN,True)
for i,(t,b) in enumerate([('Explain a classification','Use the structured detection context.'),('Interpret confidence','Clarify what the score represents.'),('Offer sorting guidance','Suggest responsible next steps.'),('Keep recent conversation','Carry the latest turns in the session.')]):
    yy=2.77+i*.73; shape(s,7.73,yy,.24,.24,[GREEN,TEAL,AMBER,BLUE][i],None,True); txt(s,t,8.15,yy-.02,3.65,.24,12,WHITE,True); txt(s,b,8.15,yy+.25,3.65,.27,9,MUTED)
line(s,7.73,5.84,12.05,5.84,'29435A',.8)
txt(s,'UI → FastAPI /api/chat → Groq → response',7.75,6.02,4.22,.30,11,WHITE,True)
txt(s,'Groq key stays server-side.',7.75,6.36,3.4,.18,9,GREEN,True)

# 08 — Architecture
s=base('A clear boundary between product and model','SYSTEM ARCHITECTURE',8)
shape(s,.68,1.60,7.30,4.96,PANEL,'223249',True)
txt(s,'WASTE ANALYSIS',.98,1.85,2.4,.24,10,GREEN,True)
node(s,1.00,2.34,1.42,'WEB UI','React / TS',BLUE)
node(s,2.85,2.34,1.42,'API','/api/analyze',TEAL)
node(s,4.70,2.34,1.42,'YOLO','adapter',AMBER)
node(s,6.48,2.34,1.38,'JSON','contract',GREEN)
line(s,2.46,2.76,2.81,2.76,TEAL,1.5); line(s,4.31,2.76,4.67,2.76,TEAL,1.5); line(s,6.16,2.76,6.44,2.76,TEAL,1.5)
txt(s,'WasteAnalysisResponse',1.02,3.46,2.24,.23,10,WHITE,True)
for i,(t,c) in enumerate([('class_name',WHITE),('category',GREEN),('confidence',TEAL),('bbox [x1,y1,x2,y2]',BLUE)]): pill(s,t,1.03+i*1.56,3.91,1.46,c)
line(s,1.03,4.60,7.54,4.60,'29435A',.8)
pill(s,'DEMO DEFAULT',1.04,4.91,1.42,AMBER,'2A2417')
txt(s,'Mock provider + local session analytics',2.72,4.93,4.57,.22,10,MUTED)
pill(s,'INTEGRATION PATH',1.04,5.51,1.70,TEAL)
txt(s,'API selector + FastAPI route + model adapter',2.98,5.53,4.24,.22,10,MUTED)
shape(s,8.30,1.60,4.15,4.96,PANEL,'223249',True)
txt(s,'AI ASSISTANT',8.62,1.85,2.4,.24,10,TEAL,True)
node(s,8.58,2.48,1.42,'CHAT','context',BLUE)
node(s,10.37,2.48,1.42,'API','/api/chat',TEAL)
node(s,9.60,3.86,1.34,'GROQ','server key',GREEN)
line(s,9.95,3.35,10.10,3.79,TEAL,1.5); line(s,11.08,3.35,10.60,3.79,TEAL,1.5)
txt(s,'Structured context, recent turns, answer',8.72,4.77,3.32,.48,11,WHITE,True,align=PP_ALIGN.CENTER)
txt(s,'Implemented routes and UI; final deployment depends on service/model configuration.',8.64,5.62,3.40,.56,9,MUTED,False,align=PP_ALIGN.CENTER)

# 09 — Impact & future
s=base('Build toward better sorting decisions','IMPACT & ROADMAP',9)
txt(s,'Potential value',.68,1.53,3.5,.30,16,WHITE,True)
vals=[('Identify','Make item types easier to recognize.'),('Explain','Show category, confidence and evidence.'),('Guide','Offer practical, contextual advice.'),('Learn','Use analysis summaries to inform next steps.')]
for i,(a,b) in enumerate(vals):
    y=2.07+i*.88; shape(s,.68,y,4.76,.68,PANEL,'223249',True); shape(s,.90,y+.19,.28,.28,[GREEN,TEAL,AMBER,BLUE][i],None,True); txt(s,a,1.38,y+.08,1.24,.24,12,WHITE,True); txt(s,b,2.66,y+.10,2.53,.40,9,MUTED)
txt(s,'Next steps',6.10,1.53,3.2,.30,16,WHITE,True)
road=[('NOW','Connect and validate model output'),('NEXT','Expand reviewed classes and guidance'),('LATER','Explore campus and municipal analytics')]
for i,(tag,t) in enumerate(road):
    y=2.12+i*1.02; pill(s,tag,6.10,y,1.05,[GREEN,TEAL,BLUE][i]); txt(s,t,7.42,y-.01,4.48,.40,12,WHITE,True)
    if i<2: line(s,6.60,y+.40,6.60,y+.98,'35536A',1.2)
shape(s,.68,6.03,11.76,.52,'132B2B',None,True)
txt(s,'DETECT',1.18,6.14,1.25,.23,12,GREEN,True,align=PP_ALIGN.CENTER); line(s,2.57,6.27,3.06,6.27,TEAL,1.3)
txt(s,'UNDERSTAND',3.13,6.14,1.56,.23,12,TEAL,True,align=PP_ALIGN.CENTER); line(s,4.88,6.27,5.36,6.27,TEAL,1.3)
txt(s,'EDUCATE',5.46,6.14,1.32,.23,12,AMBER,True,align=PP_ALIGN.CENTER); line(s,6.97,6.27,7.44,6.27,TEAL,1.3)
txt(s,'IMPROVE',7.54,6.14,1.30,.23,12,BLUE,True,align=PP_ALIGN.CENTER)
txt(s,'Potential outcomes only  ·  No physical diversion or tonnage has been measured.',.72,6.69,10,.18,9,MUTED,False,italic=True)

# 10 — Close
s=prs.slides.add_slide(blank); s.background.fill.solid(); s.background.fill.fore_color.rgb=rgb(BG)
shape(s,8.88,.76,3.60,3.60,'102634',None,True,15)
txt(s,'THE BINVISION FLOW',9.23,1.07,2.85,.24,9,GREEN,True)
for i,(lab,col) in enumerate([('UPLOAD',BLUE),('DETECT',TEAL),('SORT',GREEN)]):
    xx=9.24+i*.96
    shape(s,xx,1.74,.84,.80,PANEL,'29435A',True)
    shape(s,xx+.33,1.91,.18,.18,col,None,True)
    txt(s,lab,xx+.05,2.20,.74,.19,7,WHITE,True,align=PP_ALIGN.CENTER)
    if i<2: line(s,xx+.85,2.14,xx+.95,2.14,TEAL,1.2)
for i,(lab,col) in enumerate([('R',GREEN),('O',AMBER),('H',RED)]):
    shape(s,9.24+i*.96,3.06,.84,.39,PANEL2,None,True)
    shape(s,9.39+i*.96,3.18,.12,.12,col,None,True)
    txt(s,lab,9.62+i*.96,3.12,.22,.16,8,col,True)
txt(s,'BINVISION',.78,.80,6.9,.80,34,WHITE,True)
txt(s,'See Waste. Understand Waste.\nSort Smarter.',.82,1.80,7.3,1.00,25,WHITE,True)
txt(s,'Upload  →  Detect  →  Classify  →  Analyze  →  Ask AI',.86,3.37,8.9,.34,14,GREEN,True)
txt(s,'Turning waste images into actionable intelligence.',.84,4.29,7.9,.40,17,TEAL,True)
for i,(lab,col) in enumerate([('RECYCLABLE',GREEN),('ORGANIC',AMBER),('HAZARDOUS',RED)]): pill(s,lab,.86+i*1.84,5.12,1.63,col)
txt(s,'Team: [TEAM NAME]   ·   [MEMBER NAMES]',.86,6.26,6.1,.24,10,MUTED)
txt(s,'Demo / GitHub: [ADD LINK]     QR: [ADD IF AVAILABLE]',.86,6.61,7.6,.22,9,MUTED)
txt(s,'10 / 10',11.8,7.15,.85,.18,8,MUTED,True,align=PP_ALIGN.RIGHT)

prs.save(OUT/'BinVision_Hackathon_Pitch.pptx')

# One-page project brief
doc=Document(); sec=doc.sections[0]; sec.top_margin=DInches(.48); sec.bottom_margin=DInches(.48); sec.left_margin=DInches(.64); sec.right_margin=DInches(.64)
styles=doc.styles
styles['Normal'].font.name='Aptos'; styles['Normal'].font.size=DPt(9); styles['Normal'].font.color.rgb=DRGB(45,61,78)
styles['Normal'].paragraph_format.space_after=DPt(2)
for name,size,color in [('Title',25,(15,31,49)),('Heading 1',13,(13,112,91)),('Heading 2',10,(15,31,49))]:
    st=styles[name]; st.font.name='Aptos Display' if name=='Title' else 'Aptos'; st.font.size=DPt(size); st.font.bold=True; st.font.color.rgb=DRGB(*color)
    st.paragraph_format.space_before=DPt(4); st.paragraph_format.space_after=DPt(3)
p=doc.add_paragraph(style='Title'); p.paragraph_format.space_after=DPt(0); p.add_run('BinVision Project Brief')
p=doc.add_paragraph(); p.paragraph_format.space_after=DPt(5)
r=p.add_run('PS-1B  |  AI-Driven Waste Classification & Recycling Pipeline  |  Climate & Environment'); r.bold=True; r.font.size=DPt(9); r.font.color.rgb=DRGB(13,112,91)
p=doc.add_paragraph(); p.paragraph_format.space_after=DPt(5)
p.add_run('BinVision is a software platform that analyzes waste images, identifies detected items, assigns each item to Recyclable, Organic, or Hazardous, and presents confidence scores and bounding boxes. It helps users understand sorting decisions and provides a foundation for waste-analysis insights.')

def heading(text): doc.add_paragraph(text,style='Heading 1')
heading('The challenge')
doc.add_paragraph('Mixed materials are difficult to identify consistently. Incorrect sorting can contaminate recyclable material, while hazardous items may be handled without clear guidance. Visual feedback can help users make more informed decisions.')
heading('How the product works')
doc.add_paragraph('Upload an image, preview it, and start analysis. BinVision returns detected item names, one of three waste categories, confidence values, bounding boxes, category counts, total detections, and average confidence. Results feed the dashboard and analytics view. The BinVision AI assistant can answer questions using the current structured detection context.')
heading('Product and technical design')
doc.add_paragraph('The React and TypeScript interface uses a shared WasteAnalysisResponse contract. The analysis integration boundary is FastAPI POST /api/analyze with a YOLO model adapter and category mapping. The assistant uses FastAPI POST /api/chat and Groq. The Groq API key stays server-side.')
heading('Current demo state')
doc.add_paragraph('The frontend defaults to a mock detection provider. Its session history and analytics are local browser demo data. The standardized API route and model adapter provide the integration path; final model outputs require validation before being presented as real detection results. Assistant UI and backend chat integration are documented and use structured context when configured.')
heading('Potential value and next steps')
doc.add_paragraph('BinVision could improve waste-identification awareness, support cleaner recycling streams, and make hazardous-item guidance easier to access. Next steps include connecting and validating final model outputs, expanding reviewed waste classes, localizing disposal guidance, and exploring institutional analytics. These are potential outcomes; the project has not measured physical diversion or waste tonnage.')
heading('Team and demo')
tbl=doc.add_table(rows=2, cols=2); tbl.autofit=False; tbl.columns[0].width=DInches(1.55); tbl.columns[1].width=DInches(5.65)
for row in tbl.rows:
    row.cells[0].width=DInches(1.55); row.cells[1].width=DInches(5.65)
for i,(a,b) in enumerate([('Team','[TEAM NAME]  |  [MEMBER 1]  |  [MEMBER 2]  |  [MEMBER 3]'),('Links','Demo: [ADD URL]  |  GitHub: [ADD URL]')]):
    tbl.cell(i,0).text=a; tbl.cell(i,1).text=b
    for j,cell in enumerate(tbl.rows[i].cells):
        tcPr=cell._tc.get_or_add_tcPr(); shd=OxmlElement('w:shd'); shd.set(qn('w:fill'),'EAF3F2' if j==0 else 'F4F7FA'); tcPr.append(shd)
        for para in cell.paragraphs:
            para.paragraph_format.space_after=DPt(0)
            for run in para.runs:
                run.font.name='Aptos'; run.font.size=DPt(8); run.font.bold=(j==0); run.font.color.rgb=DRGB(13,112,91) if j==0 else DRGB(45,61,78)
doc.save(OUT/'BinVision_Project_Brief.docx')
print('Created', OUT/'BinVision_Hackathon_Pitch.pptx')
print('Created', OUT/'BinVision_Project_Brief.docx')
